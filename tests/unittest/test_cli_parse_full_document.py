"""验证整本格式从首次解析到 CLI 输出的范围不会被预估页数截断。"""

from __future__ import annotations

from doclib_test_utils import create_doclib_server

from doclib_test_utils import create_parse_services

import asyncio
import json
import re
import shlex
from collections.abc import Coroutine, Iterator
from pathlib import Path
from types import SimpleNamespace
from typing import Any, TypeVar

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from typer.testing import CliRunner

from mineru.backend import analyze
from mineru import parser
from mineru.cli.commands import parse as parse_cmd
from mineru.cli.main import app
from mineru.doclib.core.db import DatabaseManager
from mineru.doclib.core.fts import FTSManager
from mineru.doclib.services.config_svc import ConfigService
from mineru.doclib.types import (
    ContentRequestScope,
    DocContentExportRequest,
    DocContentExportResponse,
    DocContentResponse,
    ListParsesResponse,
    ParseInfo,
    ParseRequest,
    ParseResponse,
    TelemetryObservationsRequest,
)
from mineru.parser import mineru_parser
from mineru.types import Tier

_T = TypeVar("_T")


class _DoclibHarness:
    """用临时数据库连接真实解析与导出服务，避免访问常驻服务。"""

    def __init__(self, root: Path, section_count: int = 3) -> None:
        """创建多分节真实 DOCX 和独立的数据库、任务队列。"""
        self.source = root / "multi_section.docx"
        document = Document()
        for section in range(1, section_count + 1):
            if section > 1:
                document.add_section(WD_SECTION_START.NEW_PAGE)
            document.add_paragraph(f"SECTION-{section}-MARKER")
            for paragraph in range(20):
                document.add_paragraph(f"Section {section} paragraph {paragraph}: 这是测试文本，用于验证多分节解析完整性。")
        document.save(self.source)
        self.loop = asyncio.new_event_loop()
        self.db = DatabaseManager(str(root / "doclib.db"))
        self.run(self.db.initialize())
        self.service = create_parse_services(
            self.db,
            FTSManager(self.db),
            ConfigService(self.db),
            str(root / "data"),
            parse_lock_timeout_sec=1800,
        )
        self.server = create_doclib_server(
            SimpleNamespace(
                db=self.db, data_dir=str(root / "data"), parse_queue_svc=self.service.queue, ingest_svc=self.service.ingest
            )
        )
        self.output_ranges: list[str | None] = []

    def run(self, operation: Coroutine[Any, Any, _T]) -> _T:
        """在同一事件循环中执行 CLI 适配器发起的服务调用。"""
        return self.loop.run_until_complete(operation)

    async def process_pending(self) -> None:
        """执行隔离队列中的任务，保留真实 DOCX 解析与结果持久化。"""
        while (task := await self.service.queue.acquire_task()) is not None:
            assert await self.service.execution.process_doc(task)

    def close(self) -> None:
        """关闭解析使用的线程池和事件循环。"""
        self.run(self.loop.shutdown_default_executor())
        self.loop.close()


@pytest.fixture
def doclib_harness(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, request: pytest.FixtureRequest) -> Iterator[_DoclibHarness]:
    """只替换客户端传输和轮询等待，其余路径均使用真实实现。"""

    def _keep_process_logging() -> None:
        """避免后台解析把全局日志 sink 绑定到短命的 CliRunner 捕获流。"""

    monkeypatch.setattr(analyze, "configure_global_log_level", _keep_process_logging)
    monkeypatch.setattr(parser, "configure_global_log_level", _keep_process_logging)
    monkeypatch.setattr(mineru_parser, "configure_global_log_level", _keep_process_logging)
    harness = _DoclibHarness(tmp_path, section_count=getattr(request, "param", 3))

    class _Client:
        """把 CLI 的同步调用连接到隔离的异步 Doclib 服务。"""

        def __init__(self, *, timeout: int) -> None:
            """接受生产客户端超时配置，不建立网络连接。"""

        def ensure_parse(self, request: ParseRequest) -> ParseResponse:
            """使用真实入库、缓存查找和任务调度。"""
            return harness.run(
                harness.service.queue.request_parse(
                    request.path,
                    tier=request.tier,
                    page_range=request.page_range,
                    force=request.force,
                    remote=request.remote,
                )
            )

        def list_parses(self, *, ids: list[int]) -> ListParsesResponse:
            """在状态轮询前完成后台解析并读取实际持久化范围。"""
            harness.run(harness.process_pending())
            return harness.run(harness.server.list_parses(ids=ids))

        def get_doc_content(
            self,
            doc_ref: str,
            *,
            tier: Tier,
            page_range: str | None,
            after: str | None,
            limit: int,
            format: str,
            no_marker: bool,
        ) -> DocContentResponse:
            """记录范围并调用真实的内容分页和 Markdown 渲染。"""
            harness.output_ranges.append(page_range)
            return harness.run(
                harness.server.get_doc_content(
                    doc_ref,
                    tier=tier,
                    page_range=page_range,
                    after=after,
                    limit=limit,
                    format=format,
                    no_marker=no_marker,
                )
            )

        def export_doc_content(self, doc_ref: str, request: DocContentExportRequest) -> DocContentExportResponse:
            """记录范围并将真实缓存内容导出到测试目录。"""
            harness.output_ranges.append(request.page_range)
            return harness.run(harness.server.export_doc_content(doc_ref, request))

        def record_observations(self, request: TelemetryObservationsRequest) -> None:
            """隔离回归不发送使用统计。"""

    def _skip_wait(seconds: float) -> None:
        """后台任务由测试主动驱动，不需要真实轮询延迟。"""

    monkeypatch.setenv("CI", "1")
    monkeypatch.setattr(parse_cmd, "DoclibClient", _Client)
    monkeypatch.setattr(parse_cmd.time, "sleep", _skip_wait)
    try:
        yield harness
    finally:
        harness.close()


@pytest.mark.parametrize("scenario", ["first", "cache", "force", "reenter"])
@pytest.mark.parametrize("output_mode", ["stdout", "file", "json", "json_file"])
def test_multisection_docx_cli_keeps_full_document(
    doclib_harness: _DoclibHarness,
    tmp_path: Path,
    scenario: str,
    output_mode: str,
) -> None:
    """首次、缓存、重解析和异步重入均应按序保留三节全部段落。"""
    harness = doclib_harness
    runner = CliRunner()
    args = ["parse", str(harness.source), "--tier", "flash", "--no-marker"]
    if scenario in {"cache", "force"}:
        harness.run(harness.service.queue.request_parse(str(harness.source), tier="flash"))
        harness.run(harness.process_pending())
    elif scenario == "reenter":
        pending = runner.invoke(app, [*args, "--no-wait", "--json"])
        assert pending.exit_code == 0, pending.output
        assert json.loads(pending.stdout)["parse"]["status"] == "pending"
    if scenario == "force":
        args.append("--force")
    output = tmp_path / "out.md"
    if "file" in output_mode:
        args.extend(["-o", str(output)])
    if "json" in output_mode:
        args.append("--json")

    result = runner.invoke(app, args)

    assert result.exit_code == 0, result.output
    assert harness.output_ranges == [None]
    payload = json.loads(result.stdout) if "json" in output_mode else None
    if payload is not None:
        assert payload["parse"]["status"] == "done"
        assert payload["parse"]["page_range"] == "1-3"
    if "file" in output_mode:
        content = output.read_text(encoding="utf-8")
    elif payload is not None:
        content = payload["content"]["content"]
        assert payload["content"]["request_scope"]["page_range"] is None
        assert payload["content"]["truncated"] is False
    else:
        content = result.stdout
    expected = [
        text
        for section in (1, 2, 3)
        for text in [f"SECTION-{section}-MARKER", *[f"Section {section} paragraph {paragraph}:" for paragraph in range(20)]]
    ]
    assert all(content.count(text) == 1 for text in expected)
    positions = [content.index(text) for text in expected]
    assert positions == sorted(positions)


def test_multisection_docx_cli_cursor_reaches_last_section(doclib_harness: _DoclibHarness) -> None:
    """整本读取仍遵守 stdout 限额，按返回游标续读时不漏段或重复。"""
    harness = doclib_harness
    runner = CliRunner()
    args = ["parse", str(harness.source), "--tier", "flash", "--no-marker", "--limit", "200", "--json"]
    chunks: list[str] = []
    after: str | None = None
    for _ in range(100):
        result = runner.invoke(app, args + (["--after", after] if after else []))
        assert result.exit_code == 0, result.output
        payload = json.loads(result.stdout)
        assert payload["parse"]["page_range"] == "1-3"
        content = payload["content"]
        assert content["request_scope"]["limit"] == 200
        assert content["request_scope"]["page_range"] is None
        chunks.append(content["content"])
        next_request = content["next_request"]
        if next_request is None:
            assert content["truncated"] is False
            break
        assert content["truncated"] is True
        assert next_request.get("page_range") is None
        assert next_request["after"] != after
        after = next_request["after"]
    else:
        pytest.fail("DOCX 内容游标未能在有限次续读后结束")
    assert len(chunks) > 1
    combined = "\n".join(chunks)
    for section in (1, 2, 3):
        assert combined.count(f"SECTION-{section}-MARKER") == 1
        for paragraph in range(20):
            assert combined.count(f"Section {section} paragraph {paragraph}:") == 1


@pytest.mark.parametrize("doclib_harness", [11], indirect=True)
def test_docx_cli_does_not_apply_pdf_default_page_limit(doclib_harness: _DoclibHarness) -> None:
    """十一分节 DOCX 的默认 stdout 输出不能套用 PDF 前十页限制。"""
    result = CliRunner().invoke(app, ["parse", str(doclib_harness.source), "--tier", "flash", "--json"])
    assert result.exit_code == 0, result.output
    payload = json.loads(result.stdout)
    assert payload["parse"]["page_range"] == "1-11"
    assert payload["content"]["truncated"] is False
    assert payload["content"]["next_request"] is None
    assert payload["content"]["request_scope"]["page_range"] is None
    content = payload["content"]["content"]
    for section in range(1, 12):
        assert content.count(f"SECTION-{section}-MARKER") == 1
        assert content.count(f"Section {section} paragraph 19:") == 1


@pytest.mark.parametrize("doclib_harness", [11], indirect=True)
def test_docx_cli_next_commands_read_all_sections(doclib_harness: _DoclibHarness) -> None:
    """执行十一节 DOCX 实际输出的 Next 命令，验证续读不带不合法的 --pages。"""
    runner = CliRunner()
    args = ["parse", str(doclib_harness.source), "--tier", "flash", "--limit", "800"]
    chunks: list[str] = []
    commands: set[tuple[str, ...]] = set()
    for _ in range(100):
        result = runner.invoke(app, args)
        assert result.exit_code == 0, result.output
        marker = re.search(r"<!-- Next: (.*?) -->", result.stdout)
        chunks.append(result.stdout[: marker.start()] if marker else result.stdout)
        if marker is None:
            break
        command = shlex.split(marker.group(1))
        assert command[0] == "mineru"
        assert "--pages" not in command
        assert "--after" in command
        assert tuple(command) not in commands
        commands.add(tuple(command))
        args = command[1:]
    else:
        pytest.fail("DOCX 输出的 Next 命令未能读完整本")
    assert commands
    combined = "\n".join(chunks)
    for section in range(1, 12):
        assert combined.count(f"SECTION-{section}-MARKER") == 1
        assert combined.count(f"Section {section} paragraph 19:") == 1


@pytest.mark.parametrize("locator_read", [False, True])
def test_docx_explicit_logical_page_read_stays_scoped(doclib_harness: _DoclibHarness, locator_read: bool) -> None:
    """整本默认读取不改变内容接口显式范围及定位器的逻辑页选择。"""
    harness = doclib_harness
    submitted = harness.run(harness.service.queue.request_parse(str(harness.source), tier="flash"))
    harness.run(harness.process_pending())
    if locator_read:
        response = harness.run(harness.server.read_content(f"doc:{submitted.short_id}/tier:flash/page:2", no_marker=True))
    else:
        response = harness.run(harness.server.get_doc_content(submitted.sha256, tier="flash", page_range="2", no_marker=True))
    assert "SECTION-2-MARKER" in response.content
    assert "SECTION-1-MARKER" not in response.content
    assert "SECTION-3-MARKER" not in response.content
    assert response.request_scope.page_range == "2"


@pytest.mark.parametrize("cache_hit", [False, True])
@pytest.mark.parametrize("output_file", [False, True])
def test_pdf_cli_preserves_requested_range_when_batch_is_wider(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    cache_hit: bool,
    output_file: bool,
) -> None:
    """PDF 请求 2-3 页复用 1-5 页任务时，输出范围仍须保持 2-3。"""
    source = tmp_path / "document.pdf"
    source.write_bytes(b"%PDF-1.7\n")
    observed: list[str | None] = []

    class _Client:
        """模拟已完成或正在等待的更宽 PDF 缓存任务。"""

        def __init__(self, *, timeout: int) -> None:
            """接受 CLI 的客户端初始化参数。"""

        def ensure_parse(self, request: ParseRequest) -> ParseResponse:
            """保留用户范围并关联覆盖五页的已有任务。"""
            assert request.page_range == "2-3"
            return ParseResponse(
                sha256="a" * 64,
                tier="flash",
                page_range="2-3",
                status="done" if cache_hit else "pending",
                cache_hit=cache_hit,
                wait_parse_ids=[] if cache_hit else [1],
            )

        def list_parses(self, *, ids: list[int]) -> ListParsesResponse:
            """返回比当前请求范围更宽的完成任务。"""
            return ListParsesResponse(
                parses=[
                    ParseInfo(
                        id=1,
                        sha256="a" * 64,
                        short_id="aaaaaaa",
                        tier="flash",
                        page_range="1-5",
                        status="done",
                        privacy="local",
                        created_at=1,
                        updated_at=2,
                    )
                ],
                total=1,
                limit=50,
            )

        def get_doc_content(
            self,
            doc_ref: str,
            *,
            tier: Tier,
            page_range: str | None,
            after: str | None,
            limit: int,
            format: str,
            no_marker: bool,
        ) -> DocContentResponse:
            """检查 stdout 读取使用的范围与限额。"""
            observed.append(page_range)
            assert after is None
            assert limit == 30000
            return DocContentResponse(
                sha256=doc_ref,
                short_id="aaaaaaa",
                tier=tier,
                format="markdown",
                content="selected pages",
                request_scope=ContentRequestScope(page_range=page_range, limit=limit),
            )

        def export_doc_content(self, doc_ref: str, request: DocContentExportRequest) -> DocContentExportResponse:
            """检查文件导出继续携带用户指定的 PDF 范围。"""
            observed.append(request.page_range)
            Path(request.output).write_text("selected pages", encoding="utf-8")
            return DocContentExportResponse(sha256=doc_ref, short_id="aaaaaaa", tier=request.tier, output=request.output)

        def record_observations(self, request: TelemetryObservationsRequest) -> None:
            """隔离测试不发送统计。"""

    def _skip_wait(seconds: float) -> None:
        """无需为完成任务执行真实轮询延迟。"""

    monkeypatch.setenv("CI", "1")
    monkeypatch.setattr(parse_cmd, "DoclibClient", _Client)
    monkeypatch.setattr(parse_cmd.time, "sleep", _skip_wait)
    args = ["parse", str(source), "--tier", "flash", "--pages", "2-3", "--json"]
    if output_file:
        args.extend(["-o", str(tmp_path / "selected.md")])
    result = CliRunner().invoke(app, args)
    assert result.exit_code == 0, result.output
    assert json.loads(result.stdout)["parse"]["page_range"] == "2-3"
    assert observed == ["2-3"]
